"""Dissection guides (PowerPoint) -> slides -> chunks. Read only: the original .pptx files are never changed or renamed.

  extract_zip(zip_path, dest)       safe extraction (no path escapes); the zip itself is untouched
  parse_name(file name)             "Class 17 Dissection Guide - The Hip and Knee Joints.pptx" -> (17, "The Hip and Knee Joints")
  read_deck(path)                   -> [{"n", "title", "texts", "tables", "notes", "n_pictures", "n_diagrams", "hidden", "n_words"}]
                                    titles, text frames (also inside groups), tables, speaker notes; pictures are only counted
  (class map and chunking are further down in this file)

Run (repo root):  .venv/bin/python -I scripts/guides_build.py   (inventory, class map, slides, chunks)
"""
import re
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
GUIDES_DIR = ROOT / "data" / "raw" / "dissection_guides"
NAME_RE = re.compile(r"^Class\s+(\d+)\s+Dissection Guide\s*-\s*(.+?)\.pptx$", re.I)
SHORT_TEXT_WORDS = 12         # a slide with fewer words than this and at least one picture is "image-only" for retrieval


def extract_zip(zip_path, dest):
    """Extract every file below dest. A member whose path would escape dest is refused. -> list of extracted paths."""
    dest = Path(dest).resolve()
    out = []
    with zipfile.ZipFile(zip_path) as z:
        for info in z.infolist():
            target = (dest / info.filename).resolve()
            if dest != target and dest not in target.parents:
                raise ValueError(f"unsafe path in zip: {info.filename}")
            if info.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.exists():
                with z.open(info) as src, open(target, "wb") as dst:
                    dst.write(src.read())
            out.append(target)
    return out


def parse_name(name):
    m = NAME_RE.match(Path(name).name)
    return (int(m.group(1)), m.group(2).strip()) if m else (None, Path(name).stem)


def _frames(shapes):
    """Every shape of a slide, groups opened recursively."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6 and hasattr(sh, "shapes"):          # MSO_SHAPE_TYPE.GROUP
            yield from _frames(sh.shapes)


def _text_of(frame):
    lines = []
    for p in frame.paragraphs:
        t = "".join(r.text for r in p.runs).strip() or p.text.strip()
        if t:
            lines.append(("  " * min(p.level, 3)) + t)
    return lines


def read_deck(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(path))
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        title, texts, tables, pictures, diagrams = "", [], [], 0, 0
        for sh in _frames(slide.shapes):
            uri = " ".join(sh._element.xpath(".//a:graphicData/@uri")) if sh._element.tag.endswith("graphicFrame") else ""
            if "diagram" in uri or "chart" in uri:          # SmartArt / charts: their text is not readable with python-pptx
                diagrams += 1
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or (getattr(sh, "shape_type", None) == MSO_SHAPE_TYPE.PLACEHOLDER
                                                            and sh.__class__.__name__ == "PlaceholderPicture"):
                pictures += 1
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [" | ".join(c.text.strip() for c in r.cells) for r in sh.table.rows]
                tables.append("\n".join(r for r in rows if r.strip(" |")))
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                lines = _text_of(sh.text_frame)
                if sh.is_placeholder and sh.placeholder_format.type is not None and "TITLE" in str(sh.placeholder_format.type) and not title:
                    title = " ".join(l.strip() for l in lines)
                elif lines:
                    texts.append("\n".join(lines))
        if not title and slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None else ""
        words = len(re.findall(r"\w+", " ".join([title] + texts + tables)))
        hidden = slide._element.get("show") == "0"
        slides.append({"n": n, "title": re.sub(r"\s+", " ", title).strip(), "texts": texts, "tables": tables, "notes": notes,
                       "n_pictures": pictures, "n_diagrams": diagrams, "hidden": hidden, "n_words": words})
    return slides


# ----------------------------------------------------------------------------
# cleaning: what is a slide's real text
# ----------------------------------------------------------------------------
FOOTER_RE = re.compile(r"^(PAS|ANA)\d+\s*[–—-]\s*(Class \d+ Dissection Guide|Gross Human Anatomy)", re.I)
TITLE_SLIDE_RE = re.compile(r"Laboratory Session\s*[–—-]", re.I)
END_SLIDE_RE = re.compile(r"End of Laboratory Session", re.I)
SHORT_SLIDE_WORDS = 30        # a slide with fewer words is "short" and is merged into the slide before it
MERGE_MAX_CHARS = 1100


def clean_lines(slide):
    """The slide's own lines: the footer, the page number and the single letters that label a figure are dropped."""
    lines = []
    for block in slide["texts"] + slide["tables"]:
        for line in block.split("\n"):
            t = line.strip()
            if not t or FOOTER_RE.match(t) or re.fullmatch(r"\d{1,3}", t) or re.fullmatch(r"[A-Za-z]\d?", t):
                continue
            lines.append(line.rstrip())
    return lines


def slide_role(slide, lines):
    joined = " ".join(lines)
    if END_SLIDE_RE.search(joined):
        return "end"
    if slide["n"] == 1 and TITLE_SLIDE_RE.search(joined):
        return "title"
    return "content"


def slide_text(slide):
    lines = clean_lines(slide)
    text = "\n".join(lines)
    if slide.get("notes"):
        text += ("\n" if text else "") + "Speaker notes: " + slide["notes"]
    return text


def classify_slide(slide):
    """'title' | 'end' | 'text' | 'image-only' | 'empty' (for the inventory and for choosing what to index)."""
    lines = clean_lines(slide)
    role = slide_role(slide, lines)
    if role != "content":
        return role
    words = len(re.findall(r"\w+", " ".join(lines)))
    if words >= SHORT_TEXT_WORDS:
        return "text"
    if slide["n_pictures"] or slide.get("n_diagrams"):
        return "image-only"
    return "text" if words else "empty"


# ----------------------------------------------------------------------------
# class -> lecture map
# ----------------------------------------------------------------------------
import json
import math
from collections import Counter, defaultdict

from src import terms

SECTIONS_PATH = ROOT / "data" / "processed" / "sections.jsonl"
SLIDES_PATH = ROOT / "data" / "processed" / "guide_slides.jsonl"
CHUNKS_PATH = ROOT / "data" / "processed" / "guide_chunks.jsonl"
MAP_PATH = ROOT / "data" / "dissection_guides_map.json"
INVENTORY_PATH = ROOT / "reports" / "dissection_guides_inventory.md"

SYNONYMS = {"anteromedial": ["anterior", "medial"], "anterolateral": ["anterior", "lateral"], "posteromedial": ["posterior", "medial"],
            "arm": ["brachium"], "forearm": ["antebrachium"], "scapular": ["shoulder"], "deltoid": ["shoulder"],
            "oronasal": ["oral", "nasal"], "pectoral": ["axilla"]}
GENERIC = set("the and of region regions a an in to joint joints".split())
POSITION = set("anterior posterior medial lateral superior inferior".split())      # shared by many titles: they count for little
HIGH, MEDIUM, LOW = 0.65, 0.45, 0.30
CONTENT_FULL = 0.30           # a content similarity of this much counts as a full match
CONTENT_HIGH = 0.18           # below this a match cannot be "high", whatever the title says


def title_tokens(title):
    out = []
    for w in terms.tokens(title):
        for x in SYNONYMS.get(w, [w]):
            if x not in GENERIC:
                out.append(terms.stem(x))
    return out


def title_similarity(guide_title, lecture_title):
    """Weighted share of the guide title's words that the lecture title also has (positional words weigh 0.25)."""
    g = title_tokens(guide_title)
    lec = set(title_tokens(re.sub(r"^(Lecture \d+b?|ANA407 Class \d+)\s*-\s*", "", lecture_title)))
    if not g:
        return 0.0
    weight = lambda w: 0.25 if w in POSITION else 1.0
    return sum(weight(w) for w in set(g) if w in lec) / sum(weight(w) for w in set(g))


def _vec(text):
    return Counter(terms.stem(w) for w in terms.tokens(text) if len(w) > 2 and w not in GENERIC)


def _tfidf(counter, df, n_docs):
    return {w: (1 + math.log(c)) * math.log(n_docs / df[w]) for w, c in counter.items()}


def _cosine(a, b):
    num = sum(v * b.get(w, 0.0) for w, v in a.items())
    den = math.sqrt(sum(v * v for v in a.values())) * math.sqrt(sum(v * v for v in b.values()))
    return num / den if den else 0.0


def load_lectures(path=None):
    lectures, text = {}, defaultdict(list)
    for line in Path(path or SECTIONS_PATH).read_text(encoding="utf-8").splitlines():
        if line.strip():
            s = json.loads(line)
            lectures[s["lecture_id"]] = s["lecture_title"]
            text[s["lecture_id"]].append(s["text_clean"])
    return lectures, {k: " ".join(v) for k, v in text.items()}


def display_name(number, title):
    return f"Class {number}: {title}"


def build_map(decks, lectures, lecture_text, keep=None):
    """decks: {file name: slides}. -> list of map entries, one per guide file, in class order.
    keep: the previous map; an entry marked "locked": true (set by hand) is kept as it is."""
    keep = {e["file"]: e for e in (keep or []) if e.get("locked")}
    docs = {lid: _vec(t) for lid, t in lecture_text.items()}
    gdocs = {}
    for fn, slides in decks.items():
        gdocs[fn] = _vec(" ".join(" ".join(clean_lines(s)) for s in slides))
    df = Counter()
    for d in list(docs.values()) + list(gdocs.values()):
        df.update(d.keys())
    n_docs = len(docs) + len(gdocs)
    lvec = {k: _tfidf(v, df, n_docs) for k, v in docs.items()}
    numbers = Counter(parse_name(fn)[0] for fn in decks)
    entries = []
    for fn in sorted(decks, key=lambda f: (parse_name(f)[0] or 0, f)):
        number, title = parse_name(fn)
        if fn in keep:
            entries.append(keep[fn])
            continue
        g = _tfidf(gdocs[fn], df, n_docs)
        rows = []
        for lid, ltitle in lectures.items():
            t, c = title_similarity(title, ltitle), _cosine(g, lvec[lid])
            rows.append({"lecture_id": lid, "lecture": ltitle, "title_score": round(t, 2), "content_score": round(c, 2),
                         "score": round(0.5 * t + 0.5 * min(1.0, c / CONTENT_FULL), 2)})
        rows.sort(key=lambda r: -r["score"])
        best = rows[0]
        matched = [r for r in rows if r["score"] >= LOW and r["score"] >= 0.8 * best["score"]][:2] if best["score"] >= LOW else []
        conf = "unmatched" if not matched else "high" if best["score"] >= HIGH else "medium" if best["score"] >= MEDIUM else "low"
        note = ""
        if conf == "high" and best["content_score"] < CONTENT_HIGH:          # a matching title alone is not enough for "high"
            conf, note = "medium", "the title matches but the text overlaps only weakly"
        how = ("no lecture is similar enough" if not matched else
               f"title similarity {best['title_score']:.2f} + content overlap {best['content_score']:.2f}")
        entries.append({
            "class_number": number, "title": title, "display_name": display_name(number, title),
            "class_label": f"Class {number}" if numbers[number] == 1 else f"Class {number} ({title})",
            "file": fn, "n_slides": len(decks[fn]),
            "matched_lecture_id": matched[0]["lecture_id"] if matched else "", "matched_lecture": matched[0]["lecture"] if matched else "",
            "also_lecture_id": matched[1]["lecture_id"] if len(matched) > 1 else "",
            "also_lecture": matched[1]["lecture"] if len(matched) > 1 else "",
            "method": how, "score": best["score"] if matched else best["score"], "confidence": conf,
            "candidates": [{k: r[k] for k in ("lecture", "title_score", "content_score", "score")} for r in rows[:3]],
            "note": note, "locked": False})
    # Two guides on one lecture: if the other guide's text fits that lecture much better, this one is doubtful.
    by_lecture = defaultdict(list)
    for e in entries:
        if e["matched_lecture_id"] and not e["locked"]:
            by_lecture[e["matched_lecture_id"]].append(e)
    for lid, group in by_lecture.items():
        content = lambda e: e["candidates"][0]["content_score"] if e["candidates"][0]["lecture"] == e["matched_lecture"] else 0.0
        top = max(group, key=content)
        for e in group:
            if e is not top and content(e) < CONTENT_HIGH and content(top) >= 2 * max(content(e), 0.01) and e["confidence"] in ("high", "medium"):
                e["confidence"] = "low"
                e["note"] = (f"{top['display_name']} matches this lecture's text far better ({content(top):.2f} vs {content(e):.2f}); "
                             "the title matches, but the lecture's content looks like the other guide's")
    return entries


def map_by_file(path=None):
    p = Path(path or MAP_PATH)
    if not p.exists():
        return {}
    return {e["file"]: e for e in json.loads(p.read_text(encoding="utf-8"))["guides"]}


# ----------------------------------------------------------------------------
# chunks: one per slide, very short slides merged into the slide before
# ----------------------------------------------------------------------------

def unmatched_label(entry):
    """A guide that no lecture covers is granted on its own, as 'Dissection Guide, Class 25' (shown that way in the access screen)."""
    return f"Dissection Guide, {entry['class_label']}"


def unmatched_labels(path=None):
    return [unmatched_label(e) for e in map_by_file(path).values() if not e["matched_lecture"]]


def guide_label(entry):
    """Every guide can be granted on its own as 'Dissection Guide, Class 17' (matched or not)."""
    return unmatched_label(entry)


def guide_labels(path=None):
    """[(label, display name)] for all guides, in class order: the grantable guide items of the Access screen."""
    return [(guide_label(e), e["display_name"]) for e in sorted(map_by_file(path).values(), key=lambda e: (e["class_number"], e["file"]))]


def accessible_guides(allowed_lectures, path=None):
    """Map entries a user may open. allowed_lectures None = everyone (instructor). A guide opens when its matched lecture (or second
    lecture) is granted, when the guide itself is granted ('Dissection Guide, Class N'), or when the user may open EVERY lecture."""
    entries = sorted(map_by_file(path).values(), key=lambda e: (e["class_number"], e["file"]))
    if allowed_lectures is None:
        return entries
    allowed = set(allowed_lectures)
    lectures, _ = load_lectures()
    if allowed >= set(lectures.values()):
        return entries
    return [e for e in entries if e["matched_lecture"] in allowed or (e["also_lecture"] and e["also_lecture"] in allowed)
            or guide_label(e) in allowed]


def guide_chunks(decks, entries):
    by_file = {e["file"]: e for e in entries}
    out = []
    for fn in sorted(decks, key=lambda f: (parse_name(f)[0] or 0, f)):
        e = by_file[fn]
        groups = []                                # [[slide, ...]]
        for s in decks[fn]:
            lines = clean_lines(s)
            if slide_role(s, lines) != "content" or not slide_text(s).strip():
                continue
            words = len(re.findall(r"\w+", " ".join(lines)))
            if groups and words < SHORT_SLIDE_WORDS and groups[-1][-1]["n"] == s["n"] - 1 \
                    and sum(len(slide_text(x)) for x in groups[-1]) + len(slide_text(s)) <= MERGE_MAX_CHARS:
                groups[-1].append(s)
            else:
                groups.append([s])
        suffix = "" if e["class_label"] == f"Class {e['class_number']}" else "-" + re.sub(r"[^a-z0-9]+", "", e["title"].lower())[:12]
        seen_text = {}
        for grp in groups:
            first, last = grp[0]["n"], grp[-1]["n"]
            body = "\n".join(slide_text(x) for x in grp).strip()
            if len(body) > 80 and body in seen_text:           # the same step repeated in the same deck (e.g. for the other side)
                seen_text[body]["also_slides"] += [x["n"] for x in grp]
                continue
            where = f"Slide {first}" if first == last else f"Slides {first}-{last}"
            text = "\n".join(slide_text(x) for x in grp)
            out.append({
                "id": f"guide-C{e['class_number']:02d}{suffix}-s{first:03d}", "source_type": "dissection_guide", "priority": 2,
                "lecture": e["matched_lecture"] or unmatched_label(e), "lecture2": e["also_lecture"], "guide": unmatched_label(e), "matched_lecture_id": e["matched_lecture_id"],
                "match_confidence": e["confidence"], "breadcrumb": f"Dissection Guide > {e['display_name']} > {where}",
                "section_path": [e["display_name"], where], "source_url": "", "source_file": fn, "section_number": "",
                "chunk_index": 0, "n_chunks": 1, "uploaded": "", "class_number": str(e["class_number"]),
                "class_title": e["title"] if suffix else "", "display_name": e["display_name"],
                "slide": first, "slides": [x["n"] for x in grp], "also_slides": [], "text": text})
            seen_text[body] = out[-1]
    return out


def write_jsonl(path, rows):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with Path(path).open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")


def load_chunks(path=None):
    p = Path(path or CHUNKS_PATH)
    if not p.exists():
        return []
    return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]
