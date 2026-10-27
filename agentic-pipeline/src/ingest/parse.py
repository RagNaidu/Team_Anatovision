"""Parse Dr. Inglis's lecture notes into section records.

Reads  data/raw/notes/*.docx|.doc   (read-only originals: hierarchy)
       data/processed/text/*.txt    (plain-text copies: outline/script fallbacks, coverage)
Writes data/processed/sections.jsonl  (one JSON record per section)

Layouts, chosen per file:
  docx     original .docx with Heading 1-4 styles. section_path is the full chain of
           ancestor headings (topic > Heading > Heading ...). Headings are nested by
           relative level, so Heading 4 under Heading 2 (no Heading 3) works. Bullet
           nesting comes from Word list levels.
  outline  no heading styles (the .doc files): fall back to the I. / A. / 1. / a.
           numbering in the plain-text copy, plus short unmarked sub-headings.
  script   slide-by-slide video script: one section per slide.

A .doc original is converted (a copy, in a temp folder) with LibreOffice when it is
installed, to look for heading styles. Without LibreOffice, .doc files go straight to
the outline fallback.

`text` keeps the source wording (superscripts shown as 1^(st)). `text_clean` only
rewrites that markup and collapses blank lines; typos are never corrected.
Run:  python -m src.ingest.parse   (from the repo root)
"""
import json
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.text.run import Run

ROOT = Path(__file__).resolve().parents[2]
TEXT_DIR = ROOT / "data" / "processed" / "text"
RAW_DIR = ROOT / "data" / "raw" / "notes"
OUT_PATH = ROOT / "data" / "processed" / "sections.jsonl"

SOURCE_TYPE = "notes"
PRIORITY = 2
SPLIT_LIMIT = 2500   # docx sections above this are split at top-level list items
MERGE_LIMIT = 1500   # outline sections up to this many chars stay one record

BULLET_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
META_RE = re.compile(r"^(Readings|Images)\b", re.I)
HEADING_STYLE_RE = re.compile(r"^Heading\s*(\d)$")
TOPIC_STYLES = {"Title", "Title Proper"}
SLIDE_RE = re.compile(r"^Slide\s+(\d+)\s*(?:[–—-]\s*(.*))?$")
ROMAN_RE = re.compile(r"^([IVX]+)\.\s+(\S.*)$")
UPPER_RE = re.compile(r"^([A-Z])\.\s+(\S.*)$")
NUM_RE = re.compile(r"^(\d{1,2})\.\s+(\S.*)$")
LOWER_RE = re.compile(r"^([a-z])\.\s+(\S.*)$")
ALLCAPS_RE = re.compile(r"^[A-Z][A-Z0-9 ,&/\-():'.]{3,}$")
CITATION_RE = re.compile(r"\s*\(\s*(?:N\s?\d|M&D)[^)]*(?:\)|$)")
SUPERSCRIPT_CHARS = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")
SUBSCRIPT_CHARS = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")
ROMAN_VALUES = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7,
                "VIII": 8, "IX": 9, "X": 10, "XI": 11, "XII": 12}


# ----------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------

def load_lines(path):
    raw = path.read_text(encoding="utf-8-sig")
    raw = raw.replace(" ", " ").replace("\x0c", "")
    return [ln.expandtabs(4).rstrip() for ln in raw.splitlines()]


def lecture_info(stem):
    """Title and id from the file name, e.g. 'Lecture 09 - The Hand - Notes'."""
    title = re.sub(r"\s*-\s*(Student )?Notes$", "", stem, flags=re.I)
    title = re.sub(r"\s+script$", "", title, flags=re.I).strip()
    m = re.match(r"(Lecture\s+(\d+[a-z]?)|ANA407 Class (\d+))", title)
    prefix = "L" + m.group(2) if m.group(2) else "ANA407-C" + m.group(3)
    name = re.sub(r"^.*? - ", "", title)
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return title, f"{prefix}-{slug}"


def original_path(stem):
    for ext in (".docx", ".doc"):
        if (RAW_DIR / (stem + ext)).exists():
            return RAW_DIR / (stem + ext)
    return None


def clean_text(text):
    """Superscript/subscript markup to plain text, blank lines collapsed. No typo fixes."""
    text = re.sub(r"\^\(([^)]*)\)", r"\1", text)
    text = re.sub(r"(?<=\w)_\(([^)]*)\)", r"\1", text)
    text = text.replace(" ", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{2,}", "\n", text).strip()


def clean_label(text, limit=80):
    """Short heading name from a heading line that may carry a whole sentence."""
    text = CITATION_RE.sub("", text)
    text = re.split(r"\s+[–—-]\s+|:\s|\.\s", text, maxsplit=1)[0]
    text = re.sub(r"\s+", " ", text).strip(" :.-–—")
    if len(text) > limit:
        text = text[:limit].rsplit(" ", 1)[0] + "…"
    return text


def n_chars(text):
    """Non-space characters, ignoring list markers (used for coverage)."""
    text = re.sub(r"(?m)^\s*[-*]\s+", "", text)
    return len(re.sub(r"\s", "", text))


def detect_layout(name, lines):
    if "script" in name.lower() or sum(bool(SLIDE_RE.match(l)) for l in lines) >= 5:
        return "script"
    return "outline"


def finish(records_raw, stats):
    """Drop empty records, collapse blank runs, count what was dropped."""
    out = []
    for r in records_raw:
        text = re.sub(r"\n{3,}", "\n\n", "\n".join(r["lines"])).strip()
        if not text:
            stats["heading_only_dropped"] += 1
            continue
        rec = {"section_path": r["path"], "text": text}
        if r.get("readings"):
            rec["readings"] = r["readings"]
        out.append(rec)
    return out


# ----------------------------------------------------------------------------
# docx layout (Heading 1-4 styles + list levels)
# ----------------------------------------------------------------------------

def para_text(p):
    """Paragraph text with superscripts as ^(x) and subscripts as _(x), like the .txt copies."""
    parts = []
    for r in p._p.xpath(".//w:r"):
        run = Run(r, p)
        if not run.text:
            continue
        mark = "sup" if run.font.superscript else "sub" if run.font.subscript else None
        if parts and parts[-1][0] == mark:
            parts[-1][1] += run.text
        else:
            parts.append([mark, run.text])
    out = []
    for mark, t in parts:
        if mark and t.strip():
            digits = t.strip()
            if mark == "sup" and re.fullmatch(r"[0-9+\-]+", digits):
                out.append(digits.translate(SUPERSCRIPT_CHARS))  # 90⁰, as in the .txt copies
            elif mark == "sub" and re.fullmatch(r"[0-9]+", digits):
                out.append(digits.translate(SUBSCRIPT_CHARS))
            else:
                out.append(("^(%s)" if mark == "sup" else "_(%s)") % t)
        else:
            out.append(t)
    return re.sub(r"\s+", " ", "".join(out).replace(" ", " ")).strip()


def has_heading_styles(path):
    try:
        d = docx.Document(str(path))
    except Exception:
        return False
    return any(HEADING_STYLE_RE.match(p.style.name) and p.text.strip() for p in d.paragraphs)


def parse_docx(path, lecture_title, stats):
    d = docx.Document(str(path))
    lecture_name = re.sub(r"\W+", "", lecture_title.split(" - ", 1)[-1]).lower()
    topic, readings, stack, cur, recs = None, [], [], None, []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        text, style = para_text(p), p.style.name
        num = p._p.pPr.numPr if p._p.pPr is not None else None
        h = HEADING_STYLE_RE.match(style)
        if h:
            rank = int(h.group(1))
            while stack and stack[-1][0] >= rank:
                stack.pop()
            stack.append((rank, text))
            path = ([topic] if topic else []) + [label for _, label in stack]
            cur = {"path": path, "items": [], "readings": list(readings)}
            recs.append(cur)
        elif style in TOPIC_STYLES and num is None:
            if topic is None and not recs and re.sub(r"\W+", "", text).lower() == lecture_name:
                continue  # the lecture title itself (first title in the file)
            topic, stack, readings, cur = text, [], [], None
        elif num is None and META_RE.match(text):
            readings.append(text)
        else:
            if cur is None:
                cur = {"path": [topic or "Preamble"], "items": [], "readings": list(readings)}
                recs.append(cur)
            if num is not None:
                level = int(num.ilvl.val) if num.ilvl is not None else 0
                cur["items"].append((level, "  " * level + "- " + text))
            else:
                cur["items"].append((-1, text))

    out = []
    for i, r in enumerate(recs):
        if not r["items"]:
            nxt = recs[i + 1]["path"] if i + 1 < len(recs) else []
            if nxt[:len(r["path"])] == r["path"] and len(nxt) > len(r["path"]):
                stats["parent_only"] += 1  # label lives on in its children's section_path
                continue
            r["items"] = [(-1, r["path"][-1])]  # heading with no text and no children
            stats["heading_text_kept"] += 1
        out.extend(split_oversize(r, stats))
    return out


def split_oversize(rec, stats):
    """Split a long section at top-level list items; the heading chain is left intact."""
    blocks = []
    for level, line in rec["items"]:
        if level <= 0 or not blocks:
            blocks.append([line])
        else:
            blocks[-1].append(line)
    sizes = [len("\n\n".join(b)) for b in blocks]
    base = {"readings": rec["readings"]} if rec["readings"] else {}
    if sum(sizes) <= SPLIT_LIMIT or len(blocks) < 2:
        return [{"section_path": rec["path"], "text": "\n\n".join(l for _, l in rec["items"]), **base}]
    groups, cur, size = [], [], 0
    for b, sz in zip(blocks, sizes):
        if cur and size + sz > SPLIT_LIMIT:
            groups.append(cur)
            cur, size = [], 0
        cur += b
        size += sz
    groups.append(cur)
    if len(groups) < 2:
        return [{"section_path": rec["path"], "text": "\n\n".join(cur), **base}]
    stats["split_sections"] += 1
    return [{"section_path": rec["path"] + [f"part {i} of {len(groups)}"],
             "text": "\n\n".join(g), **base} for i, g in enumerate(groups, 1)]


# ----------------------------------------------------------------------------
# outline layout (no heading styles): I. / A. / 1. / a. numbering
# ----------------------------------------------------------------------------

class Node:
    def __init__(self, label, line=None, plain=False, label_only=False):
        self.label, self.plain, self.label_only = label, plain, label_only
        self.lines = [line] if line else []
        self.children = []

    def render(self, level=0):
        pad = "  " * level
        lines = [pad + l if l else l for l in self.lines]
        for c in self.children:
            lines += c.render(level + 1)
        return lines

    def own_text(self):
        return re.sub(r"\n{3,}", "\n\n", "\n".join(self.lines)).strip()

    def full_text(self):
        return re.sub(r"\n{3,}", "\n\n", "\n".join(self.render())).strip()

    def heading_only(self):
        return self.label_only and len([x for x in self.lines if x]) == 1


def marker(s, stack):
    """Return (kind, ordinal) for an outline marker line, else None."""
    m = ROMAN_RE.match(s)
    if m and m.group(1) in ROMAN_VALUES:
        last = max((e["ord"] for e in stack if e["kind"] == "roman"), default=0)
        if ROMAN_VALUES[m.group(1)] == last + 1:
            return "roman", ROMAN_VALUES[m.group(1)]
    m = UPPER_RE.match(s)
    if m:
        return "upper", ord(m.group(1)) - 64
    m = NUM_RE.match(s)
    if m:
        return "num", int(m.group(1))
    m = LOWER_RE.match(s)
    if m:
        return "lower", ord(m.group(1)) - 96
    return None


def place(stack, kind, ordinal):
    """Update the nesting stack for a new heading and return its depth."""
    if kind == "roman":
        del stack[:]
        return 0
    if ordinal > 1:
        for idx in range(len(stack) - 1, -1, -1):  # successor of an open sibling
            if stack[idx]["kind"] == kind and stack[idx]["ord"] == ordinal - 1:
                del stack[idx:]
                return idx
        for idx in range(len(stack) - 1, -1, -1):  # same kind, out of sequence
            if stack[idx]["kind"] == kind:
                del stack[idx:]
                return idx
    return len(stack)  # first item of a list (or unseen kind): child of the top


def parse_outline(lines, stats):
    tops, stack, last, saw_heading, n = [], [], None, False, len(lines)
    for i, ln in enumerate(lines):
        s = ln.strip()
        if not s:
            if last is not None and last.lines and last.lines[-1] != "":
                last.lines.append("")
            continue
        prev_blank = i == 0 or not lines[i - 1].strip()
        next_blank = i + 1 >= n or not lines[i + 1].strip()

        mk = marker(s, stack)
        plain = None
        if mk is None:
            if ALLCAPS_RE.match(s):
                plain = s
            elif prev_blank and next_blank and len(s) <= 70 and not s.endswith(".") \
                    and not BULLET_RE.match(ln):
                plain = s
            elif s.endswith(":") and len(CITATION_RE.sub("", s)) <= 60:
                plain = s
        if mk is None and plain is None:
            if last is None:
                last = Node("Preamble")
                tops.append(last)
            last.lines.append(s)
            continue
        if mk is None and not saw_heading and ALLCAPS_RE.match(s) and not tops:
            continue  # document title line, already captured in lecture_title

        kind, ordinal = mk if mk else ("plain", 0)
        while stack and stack[-1]["kind"] == "plain" and kind != "plain":
            stack.pop()  # numbered/lettered markers outrank unmarked headings
        if kind == "plain":
            if stack and stack[-1]["kind"] == "plain":
                stack.pop()  # sibling of the previous unmarked heading
        else:
            place(stack, kind, ordinal)
        saw_heading = True
        body = re.sub(r"^(?:[IVX]+|[A-Za-z]|\d{1,2})\.\s+", "", s) if mk else s
        label = (s.split()[0] + " " if mk else "") + clean_label(body)
        plain_body = re.sub(r"\s+", " ", CITATION_RE.sub("", body)).strip(" :.-–—")
        node = Node(label, s, plain=kind == "plain",
                    label_only=plain_body == clean_label(body, limit=10 ** 6) and len(plain_body) <= 80)
        (stack[-1]["node"].children if stack else tops).append(node)
        stack.append({"kind": kind, "ord": ordinal, "node": node})
        last = node

    out = []
    for top in tops:
        emit_outline(top, [], out)
    # a bare unmarked heading ("Neck (N25)") becomes a leading line of the next record
    merged, carry = [], []
    for r in out:
        if r.get("carry"):
            carry.append(r["text"])
            continue
        if carry:
            r["text"] = "\n".join(carry + [r["text"]])
            carry = []
        merged.append({"section_path": r["section_path"], "text": r["text"]})
    if carry and merged:
        merged[-1]["text"] += "\n" + "\n".join(carry)
    return merged


def emit_outline(node, prefix, out):
    """Whole node as one record when short; otherwise its own text, then its children."""
    path = prefix + [node.label]
    full = node.full_text()
    if not node.children or len(full) <= MERGE_LIMIT:
        if full:
            out.append({"section_path": path, "text": full,
                        "carry": node.plain and not node.children and node.heading_only()})
        return
    if not node.heading_only() and node.own_text():
        out.append({"section_path": path, "text": node.own_text()})
    for c in node.children:
        emit_outline(c, path, out)


# ----------------------------------------------------------------------------
# script layout
# ----------------------------------------------------------------------------

def parse_script(lines, stats):
    recs, cur = [], None
    for ln in lines:
        s = ln.strip()
        m = SLIDE_RE.match(s)
        if m:
            label = f"Slide {m.group(1)}" + (f" – {m.group(2).strip()}" if m.group(2) else "")
            cur = {"path": [label], "lines": []}
            recs.append(cur)
        elif s:
            if cur is None:
                cur = {"path": ["Preamble"], "lines": []}
                recs.append(cur)
            cur["lines"].append(s)
        elif cur is not None and cur["lines"] and cur["lines"][-1] != "":
            cur["lines"].append("")
    # a slide whose body is a single short title line ("Intro") has no spoken text
    for r in recs:
        body = [x for x in r["lines"] if x]
        if len(body) == 1 and len(body[0]) <= 40 and not body[0].endswith("."):
            r["path"] = [r["path"][0] + " – " + body[0]]
            r["lines"] = []
    return finish(recs, stats)


# ----------------------------------------------------------------------------
# .doc handling, driver
# ----------------------------------------------------------------------------

def find_soffice():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    mac_app = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return found or (str(mac_app) if mac_app.exists() else None)


def convert_doc_copy(doc_path, workdir):
    """Convert a COPY of a .doc to .docx inside workdir. Returns the .docx path or None."""
    soffice = find_soffice()
    if not soffice:
        return None
    src = Path(workdir) / "in"
    out = Path(workdir) / "out"
    src.mkdir(exist_ok=True)
    out.mkdir(exist_ok=True)
    copy = src / doc_path.name
    shutil.copy2(doc_path, copy)  # the original is never opened by soffice
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir", str(out), str(copy)],
                       check=True, capture_output=True, timeout=180)
    except (subprocess.SubprocessError, OSError):
        return None
    result = out / (doc_path.stem + ".docx")
    return result if result.exists() else None


def parse_file(path, orig=None):
    """Parse one lecture. path = its text copy (.txt); orig = the Word file to read headings from (default: the original
    in data/raw/notes with the same name)."""
    title, lecture_id = lecture_info(path.stem)
    orig = orig or original_path(path.stem)
    stats = {"heading_only_dropped": 0, "parent_only": 0, "heading_text_kept": 0,
             "split_sections": 0, "note": ""}
    lines = load_lines(path)

    docx_path = None
    with tempfile.TemporaryDirectory(prefix="anatovision_doc_") as tmp:
        if orig is not None and orig.suffix == ".docx":
            docx_path = orig
        elif orig is not None and orig.suffix == ".doc":
            converted = convert_doc_copy(orig, tmp)
            if converted is None:
                stats["note"] = ".doc not converted (LibreOffice not available)"
            docx_path = converted
        if docx_path is not None and has_heading_styles(docx_path):
            stats["layout"] = "docx"
            sections = parse_docx(docx_path, title, stats)
        else:
            stats["layout"] = detect_layout(path.name, lines)
            parser = parse_script if stats["layout"] == "script" else parse_outline
            sections = parser(lines, stats)

    records = []
    for order, sec in enumerate(sections, 1):
        rec = {
            "lecture_id": lecture_id,
            "lecture_title": title,
            "section_path": sec["section_path"],
            "breadcrumb": " > ".join([title] + [clean_text(x) for x in sec["section_path"]]),
            "text": sec["text"],
            "text_clean": clean_text(sec["text"]),
            "source_file": orig.name if orig is not None else path.name,
            "source_type": SOURCE_TYPE,
            "priority": PRIORITY,
            "order": order,
        }
        if sec.get("readings"):
            rec["readings"] = sec["readings"]
        records.append(rec)
    stats["source_chars"] = sum(n_chars(l) for l in lines)
    stats["record_chars"] = coverage_chars(records)
    return records, stats


def coverage_chars(records):
    """Source characters accounted for: record text, each distinct heading label, readings."""
    blob = re.sub(r"\W+", "", " ".join(r["text"] for r in records)).lower()
    labels = {lab for r in records for lab in r["section_path"] if not lab.startswith("part ")}
    reads = {x for r in records for x in r.get("readings", [])}
    extra = [x for x in labels | reads if re.sub(r"\W+", "", x).lower() not in blob]
    return sum(n_chars(r["text"]) for r in records) + sum(n_chars(x) for x in extra)


def coverage(stats):
    return 100 * stats["record_chars"] / max(1, stats["source_chars"])


def main():
    files = sorted(p for p in TEXT_DIR.glob("*.txt") if not p.name.startswith(".~lock"))
    all_records, per_lecture = [], []
    for path in files:
        try:
            records, stats = parse_file(path)
        except Exception as exc:  # report and continue so one odd file cannot stop the run
            print(f"!! {path.name}: {type(exc).__name__}: {exc}", file=sys.stderr)
            per_lecture.append((path, [], {"layout": "FAILED", "error": str(exc)}))
            continue
        all_records.extend(records)
        per_lecture.append((path, records, stats))

    from src import notes_upload   # uploaded notes (data/raw/notes_uploaded) replace or extend the originals
    all_records = notes_upload.apply_uploads(all_records)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with OUT_PATH.open("w", encoding="utf-8") as f:
        for r in all_records:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    report(per_lecture, all_records)


def fmt_path(rec):
    return " > ".join(rec["section_path"])


def report(per_lecture, all_records):
    print(f"Wrote {len(all_records)} sections from {len(per_lecture)} files -> {OUT_PATH.relative_to(ROOT)}\n")
    print(f"{'lecture':58} {'layout':8} {'secs':>4} {'min':>5} {'med':>5} {'max':>6} {'cov%':>5} {'multi':>5}")
    warnings = []
    for path, recs, st in per_lecture:
        title, _ = lecture_info(path.stem)
        if not recs:
            print(f"{title[:57]:58} {st['layout']:8} {0:>4}")
            warnings.append(f"{title}: produced no sections")
            continue
        sizes = [len(r["text"]) for r in recs]
        multi = sum(len(r["section_path"]) > 1 for r in recs)
        print(f"{title[:57]:58} {st['layout']:8} {len(recs):>4} {min(sizes):>5} "
              f"{int(statistics.median(sizes)):>5} {max(sizes):>6} {coverage(st):>5.0f} {multi:>5}")
        if st["layout"] != "script" and max(sizes) > SPLIT_LIMIT + 200:
            big = max(recs, key=lambda r: len(r["text"]))
            warnings.append(f"{title}: section of {max(sizes)} chars could not be split ({fmt_path(big)})")
        if len(recs) < 4:
            warnings.append(f"{title}: only {len(recs)} sections")
        if recs[0]["section_path"] == ["Preamble"] and len(recs[0]["text"]) > 1500:
            warnings.append(f"{title}: large unlabelled preamble ({len(recs[0]['text'])} chars)")
        if coverage(st) < 90:
            warnings.append(f"{title}: records cover only {coverage(st):.0f}% of source characters")

    multi_total = sum(len(r["section_path"]) > 1 for r in all_records)
    print(f"\nRecords with len(section_path) > 1: {multi_total} of {len(all_records)}")
    print("Parent-only headings kept in children's paths: "
          f"{sum(st.get('parent_only', 0) for _, _, st in per_lecture)}; "
          f"sections split for length: {sum(st.get('split_sections', 0) for _, _, st in per_lecture)}")
    shortest = min(all_records, key=lambda r: len(r["text"]))
    longest = max((r for r in all_records if "script" not in r["source_file"]), key=lambda r: len(r["text"]))
    print(f"Shortest section: {len(shortest['text'])} chars | {shortest['breadcrumb']}")
    print(f"Longest non-script section: {len(longest['text'])} chars | {longest['breadcrumb']}")
    over = [r for r in all_records if len(r["text"]) > SPLIT_LIMIT and "script" not in r["source_file"]]
    print(f"Non-script sections over {SPLIT_LIMIT} chars: {len(over)}")

    unconverted = [lecture_info(p.stem)[0] for p, _, st in per_lecture if st.get("note")]
    if unconverted:
        warnings.append(f"{len(unconverted)} .doc files not converted (LibreOffice not installed); "
                        "they use the outline-numbering fallback")
    print("\nParse warnings:")
    for w in warnings or ["none"]:
        print("  -", w)

    hand = [r for r in all_records if r["lecture_title"] == "Lecture 09 - The Hand"]
    picks = [hand[0]] + [next((r for r in hand if r["section_path"][-1] == name), None)
                         for name in ("Flexor Pollicis Brevis", "Median Nerve")]
    print(f"\n3 sample records from 'Lecture 09 - The Hand' ({len(hand)} sections total):")
    for r in [p for p in picks if p]:
        shown = {k: r[k] for k in ("lecture_id", "section_path", "breadcrumb", "text", "text_clean",
                                   "source_file", "source_type", "priority", "order", "readings") if k in r}
        for k in ("text", "text_clean"):
            if len(shown[k]) > 420:
                shown[k] = shown[k][:420] + " […truncated for display]"
        print(json.dumps(shown, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
